"""Builds the polished Team BlueBird deck from the VideoSafe original (original is never modified).

Run from anywhere:  python scripts/build_premium_deck.py
Slide 14 embeds docs/store_manager_portal.png when it exists, otherwise a placeholder frame.
"""
import copy
import io
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

DOCS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
SRC = os.path.join(DOCS, "TeamBlueBird_InventoryOptimisationAgent_Presentation_VideoSafe.pptx")
OUT = os.path.join(DOCS, "TeamBlueBird_InventoryOptimisationAgent_Presentation_Premium.pptx")
STORE_PNG = os.path.join(DOCS, "store_manager_portal.png")

NAVY, BLUE, AMBER, INK, GREY = "1F3864", "2E75B6", "F2A900", "222B3A", "5F6B7A"
PALE, LINE, MIST, WHITE, GREEN = "DEEAF6", "D9E2EE", "F3F7FC", "FFFFFF", "2E7D32"
I = 914400
LM, CW = int(0.6 * I), int(12.15 * I)  # content left edge and width (matches owner pill right edge)


def E(inches):
    return int(round(inches * I))


# ---------- drawing helpers ----------
def _shadow(shape, alpha=16000):
    sp_pr = shape._element.spPr
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:outerShdw blurRad="152400" dist="38100" dir="5400000" algn="t" rotWithShape="0">'
        f'<a:srgbClr val="1F3864"><a:alpha val="{alpha}"/></a:srgbClr></a:outerShdw></a:effectLst>'
    )
    sp_pr.append(etree.fromstring(xml))


def box(slide, x, y, w, h, fill=None, line=None, radius=0.0, shadow=False, shape=None, line_w=0.75):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, x, y, w, h)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius * I / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = int(line_w * 12700)
    else:
        s.line.fill.background()
    if shadow:
        _shadow(s)
    s.text_frame.text = ""
    return s


def tb(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, space_after=0, spc=None, line_spacing=None):
    """paras: str | list of (str | list of (str | (str, opts)))."""
    box_ = slide.shapes.add_textbox(x, y, w, h)
    tf = box_.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = int(space_after * 12700)
        if line_spacing:
            p.line_spacing = line_spacing
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            text, opts = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = text
            f = run.font
            f.name = "Arial"
            f.size = int(opts.get("size", size) * 12700)
            f.bold = opts.get("bold", bold)
            f.color.rgb = RGBColor.from_string(opts.get("color", color))
            if spc or opts.get("spc"):
                run._r.get_or_add_rPr().set("spc", str(opts.get("spc", spc)))
    return box_


def remove(shape):
    shape._element.getparent().remove(shape._element)


def by_name(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_text(shape, new):
    """Replace text keeping the first run's formatting; list = one entry per paragraph."""
    new = [new] if isinstance(new, str) else new
    for para, text in zip(shape.text_frame.paragraphs, new):
        runs = para.runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def circle_num(slide, x, y, d, n, fill=AMBER, color=NAVY, size=11):
    c = box(slide, x, y, d, d, fill=fill, line=WHITE, shape=MSO_SHAPE.OVAL, line_w=1.5)
    tb(slide, x, y, d, d, str(n), size=size, bold=True, color=color, align=PP_ALIGN.CENTER,
       anchor=MSO_ANCHOR.MIDDLE)
    return c


def stat_band(slide, y, h, label, stats):
    box(slide, LM, y, CW, h, fill=NAVY, radius=0.14, shadow=True)
    tb(slide, LM + E(0.3), y + E(0.18), E(8), E(0.26), label, size=10, bold=True, color=AMBER,
       spc=150, anchor=MSO_ANCHOR.MIDDLE)
    pad = E(0.3)
    colw = (CW - 2 * pad) // len(stats)
    for i, (big, cap) in enumerate(stats):
        cx = LM + pad + i * colw
        if i:
            box(slide, cx - E(0.18), y + E(0.55), 9525, h - E(0.75), fill="3F5F94")
        tb(slide, cx, y + E(0.5), colw - E(0.35), E(0.5), big, size=28, bold=True, color=WHITE,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(slide, cx, y + E(1.05), colw - E(0.35), h - E(1.1), cap, size=11, color="C9D6E8")


# ---------- deck-wide polish ----------
def renumber_and_rule(prs):
    for idx, slide in enumerate(prs.slides, 1):
        if idx in (1, 19):
            continue
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top >= 6300000:
                sh.text_frame.paragraphs[0].runs[0].text = str(idx)
                sh.left, sh.width = LM + CW - E(0.6), E(0.6)
                sh.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        box(slide, LM, E(6.98), CW, 9525, fill=LINE)


def fix_slide6(prs):
    s = prs.slides[5]
    ref = prs.slides[6]
    title = next(sh for sh in s.shapes if sh.name == "TextBox 22")
    title.left, title.top, title.width, title.height = 548640, 384048, 8503920, 640080
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    title.text_frame.word_wrap = True
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold = "Arial", 28 * 12700, True
        r.font.color.rgb = RGBColor.from_string(NAVY)
    for nm in ("TextBox 23", "TextBox 24"):
        remove(by_name(s, nm))
    for nm in ("Text 5", "Text 6"):
        s.shapes._spTree.append(copy.deepcopy(by_name(ref, nm)._element))


# ---------- slide 2 ----------
def slide2(prs):
    s = prs.slides[1]
    # widen cards so they end at the same right edge as the owner pill
    old_lefts = [548640 + k * 2212848 for k in range(5)]
    step_extra = (2097634 + 155448) - 2212848
    for sh in s.shapes:
        if sh.top < E(1.5) or sh.top > E(5.0):
            continue
        for k, ol in enumerate(old_lefts):
            if ol <= sh.left < ol + 2057400:
                if sh.width == 2057400:
                    sh.left, sh.width = 548640 + k * (2097634 + 155448), 2097634
                else:
                    sh.left += k * step_extra + 20117
                break
    set_text(by_name(s, "Text 11"), "Problem, use case, success criteria and delivery timeline; next steps")
    set_text(by_name(s, "Text 32"), "How a question is answered; evaluation results; live portal views")
    for nm, txt in (("Text 13", "Slides 3–5, 18–19"), ("Text 20", "Slides 6–7"), ("Text 27", "Slides 8–10"),
                    ("Text 34", "Slides 11–14"), ("Text 41", "Slides 15–17")):
        set_text(by_name(s, nm), txt)
    stat_band(s, E(5.15), E(1.65), "THE PROJECT IN NUMBERS", [
        ("17", "tools across 3 MCP servers, every query run as the signed-in person"),
        ("73 / 75", "evaluation cases passed (97.3%), rerun after every change"),
        ("0", "invented SKUs or unsupported large figures (29-answer audit)"),
        ("~10.5 s", "median time per answer at about US$0.007 per question"),
    ])


# ---------- slide 8 (text overflow) ----------
def slide8(prs):
    s = prs.slides[7]
    for nm in ("Shape 6", "Shape 10"):
        by_name(s, nm).height = E(4.05)
    for nm, h in (("Text 9", 3.0), ("Text 13", 2.7)):
        sh = by_name(s, nm)
        sh.height = E(h)
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = 13 * 12700
    arrow = by_name(s, "Shape 14")
    arrow.top = E(1.5) + E(2.025) - arrow.height // 2
    ex_top, ex_h = E(5.75), E(0.95)
    by_name(s, "Shape 15").top, by_name(s, "Shape 15").height = ex_top, ex_h
    by_name(s, "Text 17").top, by_name(s, "Text 17").height = ex_top, ex_h
    by_name(s, "Shape 16").top = ex_top + (ex_h - by_name(s, "Shape 16").height) // 2
    by_name(s, "Image 3").top = ex_top + (ex_h - by_name(s, "Image 3").height) // 2


# ---------- slide 11 ----------
def slide11(prs):
    s = prs.slides[10]
    verify = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Verify")
    cx, cy = verify.left, verify.top
    card = min((sh for sh in s.shapes if sh.shape_type == 1 and sh.left <= cx < sh.left + sh.width
                and sh.top <= cy < sh.top + sh.height and sh.width > E(2)),
               key=lambda sh: sh.width * sh.height)
    card.line.color.rgb = RGBColor.from_string(AMBER)
    card.line.width = 19050
    px = card.left + (card.width - E(1.0)) // 2
    pill = box(s, px, card.top - E(0.11), E(1.0), E(0.22), fill=AMBER, radius=0.11)
    tb(s, px, card.top - E(0.11), E(1.0), E(0.22), "KEY STEP", size=8.5, bold=True, color=NAVY,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=100)
    stat_band(s, E(5.15), E(1.7), "WHAT SITS BEHIND THE EIGHT STAGES", [
        ("17", "tested tools on 3 MCP servers, each run as the signed-in person"),
        ("≤ 10", "tool rounds per question; read-only SQL, personal details refused"),
        ("≤ 2", "rewrites if a figure or claim is unsupported, then a visible note"),
        ("8", "stages logged for audit: question, answer, tools, checks"),
    ])


# ---------- slide 12 ----------
def slide12(prs):
    s = prs.slides[11]
    remove(by_name(s, "Shape 33"))
    remove(by_name(s, "Text 34"))
    x, y, w, h = 5074920, E(5.0), 11576304 - 5074920, E(1.8)
    box(s, x, y, w, h, fill=WHITE, line=LINE, radius=0.12, shadow=True)
    box(s, x, y + E(0.18), E(0.07), h - E(0.36), fill=AMBER)
    tb(s, x + E(0.3), y + E(0.17), w - E(0.5), E(0.25),
       "DEFECT FOUND BY TESTING, THEN FIXED", size=10, bold=True, color=BLUE, spc=150,
       anchor=MSO_ANCHOR.MIDDLE)
    steps = [
        ("FOUND  ·  25/10/2026", AMBER, "The agent returned customer names, emails and phones when asked"),
        ("FIXED IN THREE PLACES", BLUE, "Request refused, queries naming those columns rejected, columns stripped from results"),
        ("LOCKED IN", GREEN, "Case T8 added and rerun with the other cases after every change"),
    ]
    gap = E(0.3)
    sw = (w - E(0.5) - E(0.2) - 2 * gap) // 3
    for i, (head, col, body) in enumerate(steps):
        sx = x + E(0.3) + i * (sw + gap)
        tb(s, sx, y + E(0.58), sw, E(0.22), head, size=9.5, bold=True, color=col, spc=60)
        tb(s, sx, y + E(0.86), sw, h - E(0.95), body, size=11, color=INK)
        if i < 2:
            box(s, sx + sw + gap // 2 - E(0.07), y + E(0.62), E(0.14), E(0.2), fill=AMBER,
                shape=MSO_SHAPE.CHEVRON)
    tb(s, 548640, E(6.3), 4206240, E(0.5),
       "The 2 misses were wording in answers about the agent itself.", size=11, color=GREY,
       align=PP_ALIGN.CENTER)


# ---------- slides 13 and 14 ----------
def portal_slide(slide, img_blob, aspect, frame_h, callouts, pins, strip_label, strip_text, cards_x_w=None):
    chrome_h = E(0.3)
    img_h = frame_h - chrome_h
    img_w = int(img_h * aspect)
    fx, fy = LM, E(1.6)
    box(slide, fx, fy, img_w, frame_h, fill=MIST, line="C9D6E8", radius=0.1, shadow=True)
    for k, col in enumerate(("E5534B", AMBER, "3DBA5D")):
        box(slide, fx + E(0.16) + k * E(0.2), fy + E(0.09), E(0.12), E(0.12), fill=col, shape=MSO_SHAPE.OVAL)
    iy = fy + chrome_h
    if img_blob:
        slide.shapes.add_picture(io.BytesIO(img_blob), fx, iy, img_w, img_h)
    else:
        ph = box(slide, fx, iy, img_w, img_h, fill=WHITE, line=BLUE, line_w=1.5)
        ph.line.dash_style = MSO_LINE.DASH
        tb(slide, fx, iy, img_w, img_h,
           ["Store Manager screenshot", "Save it as docs/store_manager_portal.png and rerun the build script"],
           size=14, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=6)
    for n, (rx, ry) in enumerate(pins, 1):
        d = E(0.3)
        circle_num(slide, fx + int(rx * img_w) - d // 2, iy + int(ry * img_h) - d // 2, d, n)
    cx = fx + img_w + E(0.35)
    cw = LM + CW - cx
    gap = E(0.14)
    ch = (frame_h - 3 * gap) // 4
    for n, (title, desc) in enumerate(callouts, 1):
        cy = fy + (n - 1) * (ch + gap)
        box(slide, cx, cy, cw, ch, fill=WHITE, line=LINE, radius=0.1, shadow=True)
        circle_num(slide, cx + E(0.18), cy + (ch - E(0.34)) // 2, E(0.34), n, size=12)
        tx = cx + E(0.7)
        tb(slide, tx, cy + E(0.1), cw - E(0.85), ch - E(0.2),
           [[(title, {"bold": True, "color": NAVY, "size": 12.5})], [(desc, {"size": 10.5, "color": GREY})]],
           anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    sy = fy + frame_h + E(0.28)
    box(slide, LM, sy, CW, E(0.82), fill=NAVY, radius=0.12, shadow=True)
    tb(slide, LM + E(0.3), sy, CW - E(0.6), E(0.82),
       [[(strip_label + "   ", {"bold": True, "color": AMBER, "size": 11, "spc": 120}),
         (strip_text, {"color": WHITE, "size": 12.5})]], anchor=MSO_ANCHOR.MIDDLE)


def slides_13_14(prs):
    s13, s14 = prs.slides[12], prs.slides[13]
    pic13 = by_name(s13, "Picture 39")
    blob13, aspect13 = pic13.image.blob, pic13.width / pic13.height
    remove(pic13)
    remove(by_name(s14, "Picture 39"))

    sub_src = by_name(prs.slides[11], "Text 1")
    for slide, sub in ((s13, "Super Manager view: the whole network, one question at a time"),
                       (s14, "Store Manager view: one store, enforced in the database")):
        el = copy.deepcopy(sub_src._element)
        slide.shapes._spTree.append(el)
        set_text(slide.shapes[-1], sub)
    title14 = copy.deepcopy(by_name(s13, "Text 0")._element)
    s14.shapes._spTree.append(title14)
    set_text(s14.shapes[-1], "Agent Portal – Store manager")

    portal_slide(
        s13, blob13, aspect13, E(4.15),
        [("Scope on every answer", "The badge and the Scope line name the stores, metric and dates"),
         ("Grounded text and tables", "Figures come from live Zava data and are checked in code first"),
         ("Private saved chats", "History is kept per person, searchable and never shared"),
         ("Role from sign-in", "Set by the access list; claiming more access changes nothing")],
        [(0.35, 0.09), (0.865, 0.26), (0.165, 0.207), (0.165, 0.95)],
        "SUPER MANAGER",
        "Sees all 8 stores: 3,997,159 units on hand across the network. The scope is stated in every answer.",
    )
    blob14 = open(STORE_PNG, "rb").read() if os.path.exists(STORE_PNG) else None
    aspect14 = 1340 / 768
    if blob14:
        from PIL import Image
        im = Image.open(io.BytesIO(blob14))
        aspect14 = im.width / im.height
    portal_slide(
        s14, blob14, aspect14, E(4.15),
        [("One store, fixed by sign-in", "This manager works in Kirkland; row-level security limits every query"),
         ("Read-only, with a reminder", "Header says live stock and sales, read-only; footer asks users to check decisions"),
         ("Four starter questions", "Excess stock, best sellers, capital tied up and seasonal pattern"),
         ("Named identity", "Name, role and store come from sign-in; chats stay private")],
        [(0.205, 0.233), (0.355, 0.058), (0.815, 0.49), (0.2, 0.95)],
        "STORE MANAGER",
        "Sees Kirkland only: 119,941 units on hand. Other stores' rows are never returned, because the limit is "
        "enforced in the database, not in the prompt.",
    )
    set_notes(s13,
              "PRESENTER: Ranjith (Implementation Lead)\n\nThis is the finished portal signed in as a Super Manager. "
              "The badge shows all 8 stores in scope, and the answer arrives as text and a table with a Scope line "
              "naming stores, metric and dates. The eight stores add up to 3,997,159 units on hand. Saved chats are "
              "private to each person, and the role comes from the access list at sign-in, not from anything the user types.")
    set_notes(s14,
              "PRESENTER: Ranjith (Implementation Lead)\n\nSame portal, signed in as a Store Manager for Kirkland. "
              "The store is fixed by the sign-in, and row-level security in the database means other stores' rows are "
              "never returned, even if the user asks. Kirkland holds 119,941 units, compared with 3,997,159 across the "
              "network for the Super Manager. The four starter questions map to our use cases: excess stock, best "
              "sellers, capital tied up and seasonal pattern. The footer reminds users to check important decisions.")


# ---------- slide 15 ----------
def slide15(prs):
    s = prs.slides[14]
    stat_band(s, E(5.2), E(1.6), "WHY THIS HOLDS UP", [
        ("7", "automated controls, each enforced in code"),
        ("1", "human approval: the agent only drafts, a manager decides"),
        ("0", "controls trusted alone: every layer is tested"),
    ])


# ---------- slide 16 ----------
def slide16(prs):
    s = prs.slides[15]
    xmap = {3520440: E(4.9), 4526280: E(6.3), 5577840: E(7.85)}
    for sh in s.shapes:
        if sh.top < 1400000:
            continue
        if sh.width == 8183879:
            sh.width = CW
        elif sh.left == 5577840:
            sh.left, sh.width = xmap[5577840], LM + CW - E(0.15) - xmap[5577840]
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        r.font.size = int(r.font.size * 1.1)
        elif sh.left in xmap:
            sh.left = xmap[sh.left]
        elif sh.left == 685800 and sh.width == 2743200:
            sh.width = E(3.9)


# ---------- slides 17 and 18 ----------
def stretch_cards(prs, idx, card_names, body_names, new_h, body_h, body_pt):
    s = prs.slides[idx]
    for nm in card_names:
        by_name(s, nm).height = new_h
    for nm in body_names:
        sh = by_name(s, nm)
        sh.height = body_h
        for p in sh.text_frame.paragraphs:
            p.space_after = int(7 * 12700)
            for r in p.runs:
                r.font.size = int(body_pt * 12700)


def main():
    prs = Presentation(SRC)
    slide2(prs)
    fix_slide6(prs)
    slide8(prs)
    slide11(prs)
    slide12(prs)
    slides_13_14(prs)
    slide15(prs)
    slide16(prs)
    stretch_cards(prs, 16, ("Shape 6", "Shape 10", "Shape 14"), ("Text 9", "Text 13", "Text 17"),
                  E(5.0), E(3.5), 15.5)
    stretch_cards(prs, 17, ("Shape 6", "Shape 10"), ("Text 9", "Text 13"), E(5.0), E(3.5), 14.5)
    renumber_and_rule(prs)
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
